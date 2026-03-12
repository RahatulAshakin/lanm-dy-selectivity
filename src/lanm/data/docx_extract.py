"""Extract a simple markdown brief from an incoming DOCX file."""

from __future__ import annotations

import re
from pathlib import Path

from lanm.filesystem import atomic_write_text
from lanm.models import BriefExtractionResult
from lanm.paths import INCOMING_DIR, PROJECT_BRIEF_PATH, REPO_ROOT

try:
    from docx import Document
except ModuleNotFoundError:  # pragma: no cover - dependency is installed in command phase
    Document = None


def _normalize_paragraph(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _iter_docx_tables(document: "Document") -> list[str]:
    rendered: list[str] = []
    for table in document.tables:
        rows = [
            [_normalize_paragraph(cell.text) for cell in row.cells]
            for row in table.rows
        ]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        padded = [row + [""] * (width - len(row)) for row in rows]
        header = padded[0]
        divider = ["---"] * width
        rendered.append("| " + " | ".join(header) + " |")
        rendered.append("| " + " | ".join(divider) + " |")
        for row in padded[1:]:
            rendered.append("| " + " | ".join(row) + " |")
        rendered.append("")
    return rendered


def _render_docx_to_markdown(docx_path: Path) -> str:
    if Document is None:  # pragma: no cover
        raise ModuleNotFoundError("python-docx")
    document = Document(docx_path)
    lines = [f"# {docx_path.stem.strip() or 'Project Brief'}", ""]
    for paragraph in document.paragraphs:
        text = _normalize_paragraph(paragraph.text)
        if not text:
            if lines and lines[-1] != "":
                lines.append("")
            continue
        style_name = paragraph.style.name.lower() if paragraph.style is not None else ""
        if style_name.startswith("heading"):
            digits = "".join(character for character in style_name if character.isdigit())
            level = int(digits) if digits else 2
            level = max(1, min(level, 6))
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)
    table_lines = _iter_docx_tables(document)
    if table_lines:
        if lines and lines[-1] != "":
            lines.append("")
        lines.extend(table_lines)
    rendered = "\n".join(lines).strip()
    return rendered + "\n"


def extract_project_brief(
    incoming_dir: Path = INCOMING_DIR,
    output_path: Path = PROJECT_BRIEF_PATH,
) -> BriefExtractionResult:
    """Extract the first incoming DOCX into markdown or write a deterministic note."""
    docx_files = sorted(
        path for path in incoming_dir.iterdir()
        if path.is_file() and path.suffix.lower() == ".docx"
    )
    if not docx_files:
        note = "No DOCX source was present under data/incoming/ during this run."
        atomic_write_text(
            output_path,
            "# Project Brief Extraction\n\n"
            f"{note}\n",
        )
        return BriefExtractionResult(
            status="missing_source",
            output_path=str(output_path.relative_to(REPO_ROOT)),
            source_path="",
            note=note,
        )
    source = docx_files[0]
    atomic_write_text(output_path, _render_docx_to_markdown(source))
    note = "Extracted from incoming DOCX."
    if len(docx_files) > 1:
        note = f"{note} Additional DOCX files were ignored after selecting {source.name}."
    return BriefExtractionResult(
        status="extracted",
        output_path=str(output_path.relative_to(REPO_ROOT)),
        source_path=str(source.relative_to(REPO_ROOT)),
        note=note,
    )

